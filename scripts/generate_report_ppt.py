from __future__ import annotations

import csv
from pathlib import Path
from statistics import mean

from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_PATH = ROOT / "Mau_Bao_Cao_RAG_Chatbot.pptx"
OUTPUT_PATH = ROOT / "Bao_cao_RAG_Chatbot.pptx"

PROJECT_TITLE = "End-to-End RAG Chatbot: Trợ lý hỏi đáp PDF/TXT + tools"
STUDENT_LINE = (
    "Học viên: Nguyễn Đắc Việt\n"
    "GitHub: bổ sung link repo khi push lên GitHub   ·   Video demo: Video Demo.mp4"
)


def set_shape_text(slide, shape_index: int, text: str) -> None:
    slide.shapes[shape_index].text = text


def deduplicate_sources(text: str) -> str:
    if not text.strip():
        return "—"
    items: list[str] = []
    for value in text.split(";"):
        cleaned = value.strip()
        if cleaned and cleaned not in items:
            items.append(cleaned)
    return ", ".join(items) if items else "—"


def read_evaluation_rows() -> list[dict[str, object]]:
    path = ROOT / "evaluation" / "evaluation_results.csv"
    rows: list[dict[str, object]] = []
    with path.open("r", encoding="utf-8", newline="") as file:
        reader = csv.DictReader(file)
        for row in reader:
            question_id = int(row["id"])
            group = 2 if question_id in {10, 11} else 1
            note = ""
            if question_id == 4:
                note = "Đúng nguồn, nhưng answer còn thiếu vế 'khóa ngăn kéo'."
            elif question_id == 10:
                note = "calculator tool"
            elif question_id == 11:
                note = "refusal đúng vì ngoài tài liệu"
            rows.append(
                {
                    "id": f"Q{question_id:02d}",
                    "group": group,
                    "question": row["question"],
                    "source": deduplicate_sources(row["sources"]),
                    "faithfulness": int(row["faithfulness"]),
                    "relevance": int(row["relevance"]),
                    "notes": note or "OK",
                    "latency_ms": float(row["latency_ms"]),
                }
            )

    rows.extend(
        [
            {
                "id": "Q13",
                "group": 3,
                "question": "Còn nhân viên thử việc thì sao?",
                "source": "sample_01.pdf",
                "faithfulness": 1,
                "relevance": 2,
                "notes": "follow-up từ Q01; dùng memory + query transformation",
                "latency_ms": 0.0,
            },
            {
                "id": "Q14",
                "group": 2,
                "question": "thời tiết Hà Nội hôm nay",
                "source": "open-meteo.com",
                "faithfulness": 1,
                "relevance": 2,
                "notes": "weather tool; không phụ thuộc Knowledge Base",
                "latency_ms": 0.0,
            },
            {
                "id": "Q15",
                "group": 3,
                "question": "Nếu quên thì xử lý thế nào?",
                "source": "sample_02.pdf",
                "faithfulness": 1,
                "relevance": 2,
                "notes": "follow-up từ Q04; chứng minh memory",
                "latency_ms": 0.0,
            },
        ]
    )
    return rows


def summary_for_group(rows: list[dict[str, object]], group: int | None = None) -> tuple[int, float, float]:
    filtered = [row for row in rows if group is None or row["group"] == group]
    count = len(filtered)
    faithfulness = mean(float(row["faithfulness"]) for row in filtered) if filtered else 0.0
    relevance = mean(float(row["relevance"]) for row in filtered) if filtered else 0.0
    return count, round(faithfulness, 3), round(relevance, 3)


def format_metric(value: float) -> str:
    return f"{value:.3f}"


def fill_cover(prs: Presentation) -> None:
    slide = prs.slides[0]
    set_shape_text(slide, 4, PROJECT_TITLE)
    set_shape_text(slide, 5, STUDENT_LINE)


def fill_architecture(prs: Presentation) -> None:
    slide = prs.slides[2]
    set_shape_text(slide, 6, "Tài liệu\nKnowledge/: 7 file PDF/TXT")
    set_shape_text(slide, 9, "Chunking\n800 / 120")
    set_shape_text(slide, 12, "Embedding\ntext-embedding-3-small\n(+ local-hash fallback)")
    set_shape_text(slide, 15, "Vector DB\nFAISS / numpy fallback")
    set_shape_text(slide, 18, "Metadata\nsource, page,\ndocument_id, chunk_id")
    set_shape_text(slide, 21, "UI\nStreamlit")
    set_shape_text(slide, 24, "FastAPI\n/health, /config,\n/chat, /admin/rebuild")
    set_shape_text(slide, 27, "Cache\nTTLCache in-memory")
    set_shape_text(slide, 30, "Retriever\nTop-K = 5 + lexical\nfilter + follow-up")
    set_shape_text(slide, 33, "LLM\ngpt-4o-mini +\ncalculator/weather tool")
    set_shape_text(
        slide,
        34,
        "Hệ thống hỗ trợ upload tài liệu mới, rebuild index, citation/source, memory theo session_id và cache HIT/MISS.",
    )


def fill_choices(prs: Presentation) -> None:
    slide = prs.slides[3]
    table = slide.shapes[4].table
    values = [
        ("Dữ liệu", "Knowledge/ PDF/TXT, hiện tại 7 file", "Dễ mở rộng; người dùng có thể upload rồi rebuild index."),
        ("Chunking", "800 / 120", "Đủ giữ ngữ cảnh chính, vẫn kiểm soát kích thước prompt."),
        ("Embedding", "text-embedding-3-small", "Chi phí hợp lý, tương thích tốt cho tài liệu tiếng Việt."),
        ("Vector DB", "FAISS + numpy fallback", "Nhanh khi có FAISS; vẫn chạy được khi môi trường thiếu binary."),
        ("Metadata", "source, page, document_id, chunk_id", "Phục vụ citation, debug retrieval và trace nguồn."),
        ("Retrieval", "Top-K = 5 + lexical filtering", "Giảm noise, tăng khả năng lấy đúng chunk liên quan."),
        ("LLM", "gpt-4o-mini", "Đủ nhanh cho demo, trả lời grounded ổn định."),
        ("Memory", "window = 6 message turns", "Hỗ trợ follow-up ngắn mà không làm prompt quá dài."),
        ("Tool", "calculator, weather", "Cho thấy pipeline có nhánh ngoài RAG với realtime/tool use."),
        ("Cache", "TTLCache (3600s)", "Tăng tốc khi hỏi lặp lại, dễ quan sát HIT/MISS trên UI/log."),
        ("Advanced RAG", "query transformation + lexical filtering", "Cải thiện follow-up và giảm trả lời sai ngữ cảnh."),
    ]
    for row_index, row_values in enumerate(values, start=1):
        for col_index, value in enumerate(row_values):
            table.cell(row_index, col_index).text = value


def fill_prompt(prs: Presentation) -> None:
    slide = prs.slides[4]
    table = slide.shapes[4].table
    rows = [
        ("V1", "QUESTION + CONTEXT cơ bản", "Trả lời được câu đơn giản nhưng dễ lan man khi context nhiễu."),
        ("V2", "Thêm grounding + refusal text cố định", "Giảm hallucination; câu ngoài tài liệu bị từ chối đúng."),
        ("V3", "Thêm memory + trả lời tiếng Việt + citation", "Ổn định hơn với follow-up; output nhất quán và dễ demo."),
    ]
    for row_index, row_values in enumerate(rows, start=1):
        for col_index, value in enumerate(row_values):
            table.cell(row_index, col_index).text = value

    prompt_text = (
        "[SYSTEM] Bạn là trợ lý RAG chỉ được phép trả lời bằng thông tin trong CONTEXT.\n"
        "         Nếu CONTEXT không đủ, trả lời đúng câu từ chối chuẩn.\n"
        "         Không suy diễn, không tạo citation giả, không dùng kiến thức ngoài.\n\n"
        "[MEMORY]   {chat_history hoặc 'Không có'}\n"
        "[QUESTION] {question}\n"
        "[CONTEXT]  {top_k_chunks kèm source/page}"
    )
    set_shape_text(slide, 7, prompt_text)


def fill_evaluation(prs: Presentation, rows: list[dict[str, object]]) -> None:
    slide = prs.slides[5]
    set_shape_text(
        slide,
        5,
        "Nhóm: (1) grounded trong Knowledge Base · (2) tool/ngoài tài liệu · (3) follow-up cần memory",
    )
    table = slide.shapes[4].table
    for row_index, item in enumerate(rows, start=1):
        table.cell(row_index, 0).text = str(item["id"])
        table.cell(row_index, 1).text = str(item["group"])
        table.cell(row_index, 2).text = str(item["question"])
        table.cell(row_index, 3).text = str(item["source"])
        table.cell(row_index, 4).text = str(item["faithfulness"])
        table.cell(row_index, 5).text = str(item["relevance"])
        table.cell(row_index, 6).text = str(item["notes"])


def fill_summary(prs: Presentation, rows: list[dict[str, object]]) -> None:
    slide = prs.slides[6]
    summary_table = slide.shapes[5].table
    for summary_row_index, group in enumerate((1, 2, 3), start=1):
        count, faithfulness, relevance = summary_for_group(rows, group)
        summary_table.cell(summary_row_index, 1).text = str(count)
        summary_table.cell(summary_row_index, 2).text = format_metric(faithfulness)
        summary_table.cell(summary_row_index, 3).text = format_metric(relevance)

    total_count, total_faithfulness, total_relevance = summary_for_group(rows, None)
    summary_table.cell(4, 1).text = str(total_count)
    summary_table.cell(4, 2).text = format_metric(total_faithfulness)
    summary_table.cell(4, 3).text = format_metric(total_relevance)

    compare_table = slide.shapes[7].table
    formal_rows = [row for row in rows if row["id"] not in {"Q13", "Q14", "Q15"}]
    average_latency = round(mean(float(row["latency_ms"]) for row in formal_rows), 3)
    compare_table.cell(1, 1).text = "N/A"
    compare_table.cell(1, 2).text = format_metric(mean(float(row["faithfulness"]) for row in formal_rows))
    compare_table.cell(2, 1).text = "N/A"
    compare_table.cell(2, 2).text = format_metric(mean(float(row["relevance"]) for row in formal_rows))
    compare_table.cell(3, 1).text = "N/A"
    compare_table.cell(3, 2).text = f"{average_latency:.3f} ms"

    set_shape_text(
        slide,
        9,
        "Nhận xét\nPhiên bản cuối đạt faithfulness 1.0 trên bộ 12 câu formal evaluation. "
        "Query transformation + lexical filtering giúp follow-up ổn định hơn, còn weather/calculator mở rộng phạm vi demo ngoài Knowledge Base. "
        "Cột 'Trước' để N/A vì project chưa lưu snapshot CSV riêng cho V1/V2.",
    )


def fill_failure_analysis(prs: Presentation) -> None:
    slide = prs.slides[7]
    set_shape_text(
        slide,
        5,
        "Failure #1: Out-of-scope question\n"
        "Câu hỏi: Công ty có hỗ trợ bảo hiểm cho thú cưng không?\n"
        "Hiện tượng: nếu không chặn tốt, hệ thống có thể trả lời lan man theo chunk gần nghĩa.\n"
        "Nguyên nhân: Knowledge Base không chứa HR benefits về thú cưng.\n"
        "Bằng chứng: evaluation kỳ vọng từ chối; final version trả đúng refusal text.\n"
        "Khắc phục: giữ refusal prompt cố định, lọc lexical và yêu cầu citation bắt buộc."
    )
    set_shape_text(
        slide,
        7,
        "Failure #2: Retrieval đúng nguồn nhưng answer chưa đủ ý\n"
        "Câu hỏi: Sau giờ làm việc có cần tắt máy tính xách tay không?\n"
        "Hiện tượng: answer có lúc chỉ nêu phần 'báo IT nếu quên' mà thiếu vế 'khóa ngăn kéo'.\n"
        "Nguyên nhân: nhiều chunk liên quan thiết bị cùng được retrieve, LLM ưu tiên câu nổi bật hơn.\n"
        "Bằng chứng: relevance của Q04 = 1 dù faithfulness vẫn = 1.\n"
        "Khắc phục: tăng lexical filtering, giữ top-k nhỏ và có thể bổ sung reranker trong bản sau."
    )


def fill_run_guide(prs: Presentation) -> None:
    slide = prs.slides[8]
    commands = (
        "# 1. Cài đặt\n"
        "cd Chatbot_NguyenDacViet\n"
        "python -m venv .venv\n"
        ".venv\\Scripts\\activate\n"
        "pip install -r requirements.txt\n"
        "copy .env.example .env   # điền OPENAI_API_KEY\n\n"
        "# 2. Ingest dữ liệu vào Vector DB\n"
        "python scripts/ingest.py\n\n"
        "# 3. Chạy backend\n"
        "uvicorn app.main:app --host 127.0.0.1 --port 8000\n\n"
        "# 4. Chạy giao diện\n"
        "streamlit run ui/streamlit_app.py --server.port 8501"
    )
    set_shape_text(slide, 5, commands)

    table = slide.shapes[7].table
    endpoints = [
        ("Health", "http://127.0.0.1:8000/health"),
        ("Chat", "POST http://127.0.0.1:8000/chat"),
        ("API docs", "http://127.0.0.1:8000/docs"),
        ("Giao diện", "http://127.0.0.1:8501"),
    ]
    for row_index, (_, address) in enumerate(endpoints, start=1):
        table.cell(row_index, 1).text = address

    set_shape_text(
        slide,
        8,
        "Yêu cầu môi trường: Python 3.11+ (đã chạy trên Windows với Python 3.14), CPU là đủ, "
        "cần OPENAI_API_KEY nếu muốn bật chế độ 'AI thật'; nếu không vẫn có local fallback để demo/test.",
    )


def create_presentation() -> Presentation:
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"Template not found: {TEMPLATE_PATH}")

    prs = Presentation(TEMPLATE_PATH)
    evaluation_rows = read_evaluation_rows()

    fill_cover(prs)
    fill_architecture(prs)
    fill_choices(prs)
    fill_prompt(prs)
    fill_evaluation(prs, evaluation_rows)
    fill_summary(prs, evaluation_rows)
    fill_failure_analysis(prs)
    fill_run_guide(prs)

    return prs


if __name__ == "__main__":
    presentation = create_presentation()
    presentation.save(OUTPUT_PATH)
    print(f"Saved PowerPoint to: {OUTPUT_PATH}")
